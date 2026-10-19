# src/report_builder.py
"""
Report Builder for SOL Research Centre.
Generates both Word (.docx) and PDF (.pdf) reports using interpreted soil data.
"""

from __future__ import annotations
import io
import os
from datetime import datetime
from docx import Document
from docx.shared import Pt
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

# ---------------------------
# WORD REPORT BUILDER
# ---------------------------
def generate_word_report(client_name: str, crop: str, interpreted_samples: list[dict], location: str = "") -> Document:
    doc = Document()

    # --- Title Page ---
    title = doc.add_paragraph()
    run = title.add_run("SOIL NUTRIENT AVAILABILITY AND SUITABILITY REPORT\n")
    run.bold = True
    run.font.size = Pt(18)
    doc.add_paragraph(f"Client: {client_name}")
    doc.add_paragraph(f"Crop: {crop}")
    if location:
        doc.add_paragraph(f"Location: {location}")
    doc.add_paragraph(f"Date: {datetime.today().strftime('%B %d, %Y')}")
    doc.add_paragraph("\n")

    # --- Summary ---
    doc.add_heading("1. Soil Laboratory Results", level=1)
    for i, sample in enumerate(interpreted_samples, start=1):
        doc.add_paragraph(f"Sample {i}")
        tbl = doc.add_table(rows=1, cols=3)
        hdr = tbl.rows[0].cells
        hdr[0].text = "Parameter"
        hdr[1].text = "Value"
        hdr[2].text = "Band"

        for param, info in sample["classified"].items():
            row = tbl.add_row().cells
            row[0].text = param
            row[1].text = str(info["value"])
            row[2].text = info["band"]

        doc.add_paragraph("\nAI Summary:")
        doc.add_paragraph(sample.get("llm_summary", "No summary generated."))

        doc.add_page_break()

    return doc

# ---------------------------
# PDF BUILDER
# ---------------------------
def generate_pdf_report(client_name: str, crop: str, interpreted_samples: list[dict], location: str = "") -> io.BytesIO:
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=A4)
    width, height = A4

    c.setFont("Helvetica-Bold", 16)
    c.drawString(50, height - 60, "SOL RESEARCH CENTRE")
    c.setFont("Helvetica", 12)
    c.drawString(50, height - 90, f"Soil Report for {client_name} - {crop}")
    if location:
        c.drawString(50, height - 110, f"Location: {location}")
    c.drawString(50, height - 130, f"Date: {datetime.today().strftime('%B %d, %Y')}")

    y = height - 160
    for i, sample in enumerate(interpreted_samples, start=1):
        c.setFont("Helvetica-Bold", 12)
        c.drawString(50, y, f"Sample {i}")
        y -= 20
        c.setFont("Helvetica", 10)
        for param, info in sample["classified"].items():
            text = f"{param}: {info['value']} ({info['band']})"
            c.drawString(60, y, text)
            y -= 15
            if y < 100:
                c.showPage()
                y = height - 100
        y -= 10
        c.setFont("Helvetica-Oblique", 10)
        summary = sample.get("llm_summary", "No summary generated.")
        for line in summary.splitlines():
            c.drawString(70, y, line)
            y -= 12
            if y < 80:
                c.showPage()
                y = height - 100
        c.showPage()

    c.save()
    buffer.seek(0)
    return buffer
